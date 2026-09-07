"""Safe text extraction for ProofLearn AI assignment uploads."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
import re
import zipfile

import pymupdf
from docx import Document

from modules.config import MAX_UPLOAD_SIZE_MB, SUPPORTED_FILE_TYPES


class DocumentExtractionError(ValueError):
    """Raised when an uploaded document cannot be safely converted to text."""


@dataclass(frozen=True)
class ExtractedDocument:
    """Text and basic provenance returned by the extraction engine."""

    filename: str
    file_type: str
    text: str
    character_count: int
    word_count: int
    paragraph_count: int
    page_count: int | None = None


def _normalise_text(text: str) -> str:
    """Remove extraction noise while preserving paragraph boundaries."""
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
    cleaned = "\n".join(lines)
    return re.sub(r"\n{3,}", "\n\n", cleaned).strip()


def _extract_txt(data: bytes) -> tuple[str, None]:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding), None
        except UnicodeDecodeError:
            continue
    raise DocumentExtractionError("The text file uses an unsupported encoding.")


def _extract_docx(data: bytes) -> tuple[str, None]:
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            members = archive.infolist()
            total_uncompressed = sum(member.file_size for member in members)
            total_compressed = max(1, sum(member.compress_size for member in members))
            if len(members) > 2_000 or total_uncompressed > 50 * 1024 * 1024:
                raise DocumentExtractionError("The Word file expands beyond the safe processing limit.")
            if total_uncompressed / total_compressed > 200:
                raise DocumentExtractionError("The Word file has an unsafe compression ratio.")
        document = Document(BytesIO(data))
        blocks: list[str] = []
        blocks.extend(paragraph.text for paragraph in document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells]
                blocks.append(" | ".join(value for value in values if value))
        return "\n\n".join(blocks), None
    except DocumentExtractionError:
        raise
    except Exception as exc:
        raise DocumentExtractionError(
            "The Word file could not be read. It may be damaged or password protected."
        ) from exc


def _extract_pdf(data: bytes) -> tuple[str, int]:
    try:
        with pymupdf.open(stream=data, filetype="pdf") as document:
            if document.needs_pass:
                raise DocumentExtractionError(
                    "Password protected PDF files are not supported."
                )
            pages = [page.get_text("text") for page in document]
            return "\n\n".join(pages), document.page_count
    except DocumentExtractionError:
        raise
    except Exception as exc:
        raise DocumentExtractionError(
            "The PDF could not be read. It may be damaged or use an unsupported format."
        ) from exc


def extract_document(data: bytes, filename: str) -> ExtractedDocument:
    """Extract text from a validated PDF, DOCX or TXT upload.

    The filename extension selects the parser. Content is kept in memory and is
    not written to disk, which limits storage of student submissions in the MVP.
    """
    if not filename or "." not in filename:
        raise DocumentExtractionError("The uploaded file must have a valid filename.")
    if not data:
        raise DocumentExtractionError("The uploaded file is empty.")

    max_bytes = MAX_UPLOAD_SIZE_MB * 1024 * 1024
    if len(data) > max_bytes:
        raise DocumentExtractionError(
            f"The file is larger than the {MAX_UPLOAD_SIZE_MB} MB upload limit."
        )

    file_type = Path(filename).suffix.lower().lstrip(".")
    if file_type not in SUPPORTED_FILE_TYPES:
        supported = ", ".join(extension.upper() for extension in SUPPORTED_FILE_TYPES)
        raise DocumentExtractionError(f"Unsupported file type. Use {supported}.")

    extractors = {
        "txt": _extract_txt,
        "docx": _extract_docx,
        "pdf": _extract_pdf,
    }
    raw_text, page_count = extractors[file_type](data)
    text = _normalise_text(raw_text)

    if not text:
        message = "No readable text was found in this document."
        if file_type == "pdf":
            message += " Scanned PDFs will require OCR support in a later milestone."
        raise DocumentExtractionError(message)

    paragraphs = [part for part in text.split("\n\n") if part.strip()]
    return ExtractedDocument(
        filename=Path(filename).name,
        file_type=file_type,
        text=text,
        character_count=len(text),
        word_count=len(re.findall(r"\b[A-Za-z]+(?:['’-][A-Za-z]+)?\b", text)),
        paragraph_count=len(paragraphs),
        page_count=page_count,
    )
