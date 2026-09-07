from io import BytesIO
import zipfile

import pymupdf
import pytest
from docx import Document

from modules.document_reader import DocumentExtractionError, extract_document


def _docx_bytes() -> bytes:
    stream = BytesIO()
    document = Document()
    document.add_heading("Student Assignment", level=1)
    document.add_paragraph("Technology supports learning when it is used responsibly.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Evidence"
    table.cell(0, 1).text = "Reflection"
    document.save(stream)
    return stream.getvalue()


def _pdf_bytes() -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "This is a readable student assignment.")
    data = document.tobytes()
    document.close()
    return data


def test_extracts_utf8_text() -> None:
    result = extract_document("First paragraph.\n\nSecond paragraph.".encode(), "work.txt")
    assert result.file_type == "txt"
    assert result.word_count == 4
    assert result.paragraph_count == 2


def test_extracts_docx_paragraphs_and_tables() -> None:
    result = extract_document(_docx_bytes(), "assignment.docx")
    assert "Technology supports learning" in result.text
    assert "Evidence | Reflection" in result.text
    assert result.paragraph_count == 3
    assert result.page_count is None


def test_extracts_pdf_and_page_count() -> None:
    result = extract_document(_pdf_bytes(), "assignment.pdf")
    assert "readable student assignment" in result.text
    assert result.page_count == 1


@pytest.mark.parametrize("filename", ["work.doc", "work.xlsx", "work.exe"])
def test_rejects_unsupported_file_types(filename: str) -> None:
    with pytest.raises(DocumentExtractionError, match="Unsupported file type"):
        extract_document(b"content", filename)


def test_rejects_empty_upload() -> None:
    with pytest.raises(DocumentExtractionError, match="empty"):
        extract_document(b"", "assignment.txt")


def test_rejects_textless_pdf_with_clear_ocr_message() -> None:
    document = pymupdf.open()
    document.new_page()
    data = document.tobytes()
    document.close()
    with pytest.raises(DocumentExtractionError, match="OCR"):
        extract_document(data, "scan.pdf")


def test_rejects_damaged_docx() -> None:
    with pytest.raises(DocumentExtractionError, match="could not be read"):
        extract_document(b"not a word document", "damaged.docx")


def test_rejects_docx_with_unsafe_compression_ratio() -> None:
    stream = BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "A" * 2_000_000)
    with pytest.raises(DocumentExtractionError, match="unsafe compression ratio"):
        extract_document(stream.getvalue(), "unsafe.docx")
